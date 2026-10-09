import copy
import json
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from unittest.mock import patch

import openpyxl

from music_eval.engine import evaluate_row, last_item, load_records, make_template, pair_stats, run
from music_eval.judge import JudgeError, ModelConfig, ModelJudge, validate_answer


ANSWER = {
    "intent": "search_only", "category": "scenario", "requires_search": True,
    "nlg_status": "pass", "nlg_reason": "回复列出搜索候选，并符合放松需求。",
    "needs_review": False, "review_reason": "", "issues": [],
    "search_quality": {"logic": 4, "diversity": None, "timeliness": None, "reason": "符合需求；多样性和时效性不适用。"},
    "play_quality": {"logic": None, "diversity": None, "timeliness": None, "reason": "无播放列表。"},
}


class FakeJudge:
    config = ModelConfig("https://example.com/v1", "test-model")

    def __init__(self, answer=None):
        self.answer = copy.deepcopy(answer or ANSWER)
        self.payloads = []

    def evaluate(self, payload):
        self.payloads.append(payload)
        return copy.deepcopy(self.answer)


def row(**kwargs):
    base = {"Query": "推荐适合放松的歌", "维度": "场景搜歌-只搜不播",
            "实际落域": "【['闲聊.其他', '多媒体.音乐.推理']】", "通道类型": "【['快', '慢']】",
            "music_search入参_query": "放松舒缓的歌", "music_search响应_candidates": "A - 歌手甲\nB - 歌手乙",
            "asr新": "推荐A和B", "标注备注": "不要让人工答案泄漏给模型"}
    base.update(kwargs)
    return base


class RuleTests(unittest.TestCase):
    def test_last_route_and_channel(self):
        self.assertEqual(last_item("【['快', '快', '慢']】"), "慢")
        self.assertEqual(last_item("【['慢', '慢']】【['快']】"), "快")
        r = evaluate_row(row(), [], FakeJudge())
        self.assertEqual(r["自动_最终落域"], "多媒体.音乐.推理")
        self.assertEqual(r["自动_PASS_FAIL"], "PASS")

    def test_search_only_missing_play_is_not_failure(self):
        r = evaluate_row(row(**{"出歌单信息": "未找到 mediaList"}), [], FakeJudge())
        self.assertEqual(r["自动_PASS_FAIL"], "PASS")

    def test_search_play_missing_play_is_failure(self):
        answer = copy.deepcopy(ANSWER)
        answer["intent"] = "search_play"
        r = evaluate_row(row(**{"维度": "又搜又播"}), [], FakeJudge(answer))
        self.assertEqual(r["自动_PASS_FAIL"], "FAIL")
        self.assertIn("调用工具错误", r["自动_问题标签"])
        self.assertEqual(r["自动_搜索逻辑分"], 4)

    def test_tool_error_independent_of_final_pass(self):
        r = evaluate_row(row(**{"music_search入参_query": ""}), [], FakeJudge())
        self.assertEqual(r["自动_PASS_FAIL"], "PASS")
        self.assertIn("调用工具错误", r["自动_问题标签"])

    def test_followup_play_does_not_need_search(self):
        answer = copy.deepcopy(ANSWER)
        answer.update(intent="followup_play", requires_search=False)
        r = evaluate_row(row(**{"预期行为": "先搜后播", "music_search入参_query": "",
                                "歌单入参_playlist": "A - 歌手甲", "出歌单信息": "A - 歌手甲"}),
                         [row()], FakeJudge(answer))
        self.assertEqual(r["自动_PASS_FAIL"], "PASS")
        self.assertNotIn("调用工具错误", r["自动_问题标签"])

    def test_no_model_no_fake_pass_or_scores(self):
        r = evaluate_row(row(), [])
        self.assertEqual(r["自动_PASS_FAIL"], "")
        self.assertEqual(r["自动_搜索逻辑分"], "")
        self.assertEqual(r["自动_复核状态"], "待复核")

    def test_merged_rounds_not_scored(self):
        judge = FakeJudge()
        r = evaluate_row(row(**{"歌单入参_playlist": "第1轮:A\n第2轮:B"}), [], judge)
        self.assertEqual(r["自动_PASS_FAIL"], "")
        self.assertEqual(len(judge.payloads), 0)

    def test_asr_source_and_human_labels_not_sent(self):
        judge = FakeJudge()
        evaluate_row(row(), [], judge)
        self.assertEqual(judge.payloads[0]["nlg_source"], "asr")
        self.assertNotIn("标注备注", judge.payloads[0]["record"])
        judge = FakeJudge()
        evaluate_row(row(**{"原始NLG文本": "原始回复"}), [], judge)
        self.assertEqual(judge.payloads[0]["nlg"], "原始回复")

    def test_explicit_search_only_unwanted_play(self):
        r = evaluate_row(row(**{"歌单入参_playlist": "A", "出歌单信息": "A"}), [], FakeJudge())
        self.assertEqual(r["自动_PASS_FAIL"], "FAIL")

    def test_evidence_pending(self):
        answer = copy.deepcopy(ANSWER)
        answer.update(nlg_status="pending", needs_review=True, review_reason="电影归属缺参考证据")
        r = evaluate_row(row(), [], FakeJudge(answer))
        self.assertEqual(r["自动_PASS_FAIL"], "")
        self.assertIn("电影归属", r["自动_复核原因"])

    def test_keyword_phrases_not_automatically_wrong(self):
        r = evaluate_row(row(**{"music_search入参_query": "放松 舒缓 钢琴"}), [], FakeJudge())
        self.assertNotIn("传参错误", r["自动_问题标签"])

    def test_gateway_failure_pending_but_objective_fail_retained(self):
        judge = FakeJudge()
        judge.evaluate = lambda p: (_ for _ in ()).throw(JudgeError("接口HTTP401"))
        r = evaluate_row(row(**{"维度": "又搜又播"}), [], judge)
        self.assertEqual(r["自动_PASS_FAIL"], "FAIL")
        self.assertIn("401", r["自动_复核原因"])

    def test_pair_jaccard_is_descriptive(self):
        rows = [("s", 2, row(**{"配对组ID": "p", "配对类型": "同诉求重复"})),
                ("s", 3, row(**{"配对组ID": "p", "配对类型": "同诉求重复", "music_search响应_candidates": "B - 歌手乙\nC - 歌手丙"}))]
        result = pair_stats(rows)
        self.assertEqual(result[0]["jaccard"], 0.3333)


class FileTests(unittest.TestCase):
    def test_export_preserves_input_and_html_escapes(self):
        with tempfile.TemporaryDirectory() as temp:
            temp = Path(temp)
            input_ = temp / "input.xlsx"
            w = openpyxl.Workbook()
            w.active.append(["Query", "维度", "asr新", "是否通过/存在问题"])
            w.active.append(["<script>alert(1)</script>", "只搜不播", "推荐A", "人工PASS"])
            w.save(input_)
            original = input_.read_bytes()
            output, report, meta = run(input_, temp / "out", FakeJudge())
            self.assertEqual(original, input_.read_bytes())
            out = openpyxl.load_workbook(output)
            self.assertEqual(out.active.cell(2,4).value, "人工PASS")
            self.assertNotIn("<script>alert(1)</script>", report.read_text(encoding="utf-8"))
            self.assertIn("&lt;script&gt;", report.read_text(encoding="utf-8"))
            self.assertEqual(json.loads(meta.read_text())["processed"], 1)
            out.close()

    def test_history_sorted_by_round_not_row_order(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "input.xlsx"
            w = openpyxl.Workbook()
            w.active.append(["Query", "会话ID", "轮次"])
            w.active.append(["播放", "A", 2])
            w.active.append(["推荐", "A", 1])
            w.save(path)
            judge = FakeJudge()
            run(path, Path(temp) / "out", judge)
            self.assertEqual(judge.payloads[0]["history"][0]["Query"], "推荐")

    def test_duplicate_round_rejected(self):
        w = openpyxl.Workbook()
        w.active.append(["Query", "会话ID", "轮次"])
        w.active.append(["A", "group", 1])
        w.active.append(["B", "group", 1])
        with self.assertRaises(ValueError):
            load_records(w)

    def test_template_has_expected_columns(self):
        with tempfile.TemporaryDirectory() as temp:
            p = Path(temp) / "template.xlsx"
            make_template(p)
            w = openpyxl.load_workbook(p)
            self.assertIn("原始NLG文本", [c.value for c in w.active[1]])
            self.assertEqual(len(load_records(w)), 2)
            w.close()

    def test_cancel_exports_partial_with_counts(self):
        with tempfile.TemporaryDirectory() as temp:
            p = Path(temp) / "template.xlsx"
            make_template(p)
            calls = []
            outputs = run(p, temp, progress=lambda d,t: calls.append(d), cancelled=lambda: bool(calls))
            meta = json.loads(outputs[2].read_text())
            self.assertEqual(meta["processed"], 1)
            self.assertEqual(meta["input_records"], 2)
            self.assertTrue(meta["cancelled"])


class GatewayTests(unittest.TestCase):
    def test_schema_rejects_bad_scores_and_boolean_scores(self):
        for score in (6, -1, True, "5"):
            answer = copy.deepcopy(ANSWER)
            answer["search_quality"]["logic"] = score
            with self.assertRaises(JudgeError):
                validate_answer(answer)

    def test_https_required_and_suffix(self):
        with self.assertRaises(ValueError):
            ModelConfig("http://company.example/v1", "m").endpoint()
        self.assertEqual(ModelConfig("https://company.example/custom/v1", "m").endpoint(),
                         "https://company.example/custom/v1/chat/completions")

    def test_real_http_compatible_gateway_request(self):
        received = []
        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass
            def do_POST(self):
                received.append((self.path, json.loads(self.rfile.read(int(self.headers["Content-Length"])))))
                response = json.dumps({"choices": [{"message": {"content": json.dumps(ANSWER)}}]}).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(response)
        server = HTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            judge = ModelJudge(ModelConfig(f"http://127.0.0.1:{server.server_port}/v1", "mock-model"))
            result = judge.evaluate({"record": {"Query": "推荐放松歌曲"}})
            self.assertEqual(result["nlg_status"], "pass")
            self.assertEqual(received[0][0], "/v1/chat/completions")
            self.assertEqual(received[0][1]["model"], "mock-model")
        finally:
            server.shutdown()
            server.server_close()
            thread.join()


if __name__ == "__main__":
    unittest.main()
