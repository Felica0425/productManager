import ast
import hashlib
import html
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .judge import JudgeError, PROMPT_VERSION, SYSTEM_PROMPT

FIELDS = ["用例ID", "会话ID", "轮次", "维度", "Query", "预期行为", "预期需搜索",
          "实际落域", "通道类型", "music_search入参_query", "music_search响应_candidates",
          "歌单入参_playlist", "出歌单信息", "原始NLG文本", "asr新", "评测日期",
          "参考资料", "时效窗口", "配对组ID", "配对类型"]
OUTPUT_FIELDS = ["自动_PASS_FAIL", "自动_复核状态", "自动_行为分类", "自动_音乐分类", "自动_最终落域", "自动_最终通道",
                 "自动_问题标签", "自动_问题证据", "自动_结果理由", "自动_复核原因",
                 "自动_搜索逻辑分", "自动_搜索多样性分", "自动_搜索时效分", "自动_搜索评分理由",
                 "自动_播放逻辑分", "自动_播放多样性分", "自动_播放时效分", "自动_播放评分理由",
                 "自动_运行ID", "自动_评分规则版本"]
INTENT_NAMES = {"search_only": "只搜不播", "search_play": "又搜又播", "followup_play": "先搜后播（播放轮）",
                "unknown": "待确认", "other": "其他"}


def text(value):
    return "" if value is None else str(value).strip()


def present(value):
    return text(value).strip('"【】 ') not in {"", "[]", "{}", "null", "None", "未找到 mediaList", "未找到mediaList"}


def last_item(value):
    """Accept nested/exported Python lists, but never execute spreadsheet text."""
    raw = text(value)
    parts = re.findall(r"\[[^\[\]]*\]", raw)
    if parts:
        try:
            items = ast.literal_eval(parts[-1])
            return text(items[-1]) if items else ""
        except (ValueError, SyntaxError, TypeError):
            return ""
    return raw.strip("【】 ")


def song_list(value):
    raw = text(value).strip('"')
    if not present(raw):
        return []
    # Preserve versions and artists. Exact text overlap is descriptive, not a correctness score.
    try:
        parsed = json.loads(raw)
        if isinstance(parsed, list):
            return [json.dumps(v, ensure_ascii=False, sort_keys=True) if isinstance(v, dict) else text(v) for v in parsed if v]
    except ValueError:
        pass
    return [re.sub(r"\s+", " ", s).strip() for s in raw.splitlines() if s.strip()]


def merged_rounds(row):
    return any(re.search(r"第\s*[12一二]\s*轮\s*[:：]", text(row.get(k))) for k in
               ("music_search入参_query", "music_search响应_candidates", "歌单入参_playlist"))


def explicit_intent(row):
    expected = text(row.get("预期行为"))
    if expected:
        mapping = {v: k for k, v in INTENT_NAMES.items()}
        mapping.update({"先搜后播": "followup_play", "播放轮": "followup_play", **{k: k for k in INTENT_NAMES}})
        return mapping.get(expected, "unknown")
    dimension = text(row.get("维度"))
    if "只搜不播" in dimension:
        return "search_only"
    if "又搜又播" in dimension:
        return "search_play"
    if "先搜后播" in dimension:
        return "followup_play"
    query = text(row.get("Query"))
    if text(row.get("会话ID")) and text(row.get("轮次")) not in {"", "1"}:
        # Context determines whether this is inherited playback or a fresh search.
        return "unknown"
    if re.search(r"别播|不要播|不播放|只.*推荐|只.*找", query):
        return "search_only"
    if re.search(r"播放|帮我放|给我放|放点|放一些|播点|播一些", query):
        return "search_play"
    return "unknown"


def parse_need_search(row):
    v = text(row.get("预期需搜索")).lower()
    if v in {"是", "true", "1", "yes"}:
        return True
    if v in {"否", "false", "0", "no"}:
        return False
    return None


def evaluate_row(row, history, judge=None):
    result = {field: "" for field in OUTPUT_FIELDS}
    result["自动_最终落域"] = last_item(row.get("实际落域"))
    result["自动_最终通道"] = last_item(row.get("通道类型"))
    reviews, failures, issues = [], [], []
    if merged_rounds(row):
        result.update({"自动_复核状态": "待复核", "自动_行为分类": "多轮合并，未评分",
                       "自动_复核原因": "多轮工具记录合并在同一行，无法可靠对应当前轮播放结果；请每轮一行。",
                       "自动_结果理由": "未将上轮结果视为末轮结果。"})
        return result
    intent = explicit_intent(row)
    nlg = text(row.get("原始NLG文本")) or text(row.get("asr新"))
    candidates = song_list(row.get("music_search响应_candidates"))
    played = song_list(row.get("出歌单信息"))
    payload = {
        "record": {k: text(row.get(k)) for k in FIELDS},
        "history": [{k: text(r.get(k)) for k in FIELDS} for r in history],
        "nlg_source": "original" if text(row.get("原始NLG文本")) else "asr",
        "nlg": nlg, "final_domain": result["自动_最终落域"], "final_channel": result["自动_最终通道"],
        "search_candidates": candidates, "played_tracks": played,
        "known_intent": intent,
    }
    answer = None
    if judge:
        try:
            answer = judge.evaluate(payload)
        except JudgeError as e:
            reviews.append(str(e))
    else:
        reviews.append("未连接模型：回复准确性、隐含意图及歌曲质量尚未评估。")
    requires_search = parse_need_search(row)
    if answer:
        if intent == "unknown":
            intent = answer["intent"]
        elif answer["intent"] != intent:
            reviews.append("模型识别的行为与预期/维度分类不一致，请核对。")
        if requires_search is None:
            requires_search = answer["requires_search"]
        result["自动_结果理由"] = answer["nlg_reason"]
        if answer["nlg_status"] == "fail":
            failures.append("最终回复不满足需求：" + answer["nlg_reason"])
        if answer["nlg_status"] == "pending" or answer["needs_review"]:
            reviews.append(answer["review_reason"] or "模型判断所需证据不足。")
        issues.extend(answer["issues"])
        result["自动_音乐分类"] = {"scenario": "场景搜歌", "reasoning": "推理音乐", "unknown": "待确认"}[answer["category"]]
        for key, prefix, tracks in (("search_quality", "搜索", candidates), ("play_quality", "播放", played)):
            q = answer[key]
            for dim, label in (("logic", "逻辑"), ("diversity", "多样性"), ("timeliness", "时效")):
                result[f"自动_{prefix}{label}分"] = q[dim] if tracks else ""
            result[f"自动_{prefix}评分理由"] = q["reason"] if tracks else "无本轮列表，未评分。"
    if requires_search is None and intent in {"search_only", "search_play"}:
        requires_search = True
    if intent == "followup_play" and not history:
        reviews.append("继承播放缺少同会话前轮记录，无法核验继承是否正确。")
    if intent == "unknown":
        reviews.append("无法确定播放意图；请填写预期行为。")
    if not nlg:
        # Missing an exported reply is missing evidence, not proof the assistant stayed silent.
        reviews.append("缺少原始NLG和asr新，无法核验最终回复。")
    if requires_search and not present(row.get("music_search入参_query")):
        issues.append({"label": "调用工具错误", "reason": "需要搜索，但music_search入参_query为空（按约定视为未调用）。"})
    if intent in {"search_play", "followup_play"} and result["自动_最终落域"].startswith("闲聊"):
        issues.append({"label": "落域错误", "reason": "音乐播放意图的最终落域为" + result["自动_最终落域"] + "。"})
    if "闲聊" in text(row.get("维度")) and intent in {"search_only", "search_play", "followup_play"} and result["自动_最终通道"] == "快":
        issues.append({"label": "落域错误", "reason": "维度标注为闲聊混合音乐诉求，但最终通道为快，预期慢。"})
    if intent in {"search_play", "followup_play"}:
        if not present(row.get("歌单入参_playlist")):
            issues.append({"label": "调用工具错误", "reason": "需要播放，但歌单入参_playlist为空（按约定视为未调用）。"})
        if not played:
            failures.append("有播放意图，但出歌单信息为空或未找到mediaList，播放失败。")
    elif intent == "search_only" and (present(row.get("歌单入参_playlist")) or played):
        failures.append("预期只搜不播，但出现播放调用/播放结果，行为不满足需求。")
        issues.append({"label": "调用工具错误", "reason": "只搜不播场景发生播放。"})
    if candidates and len(candidates) != len(set(candidates)):
        reviews.append("搜索列表存在完全重复条目，请结合列表大小复核多样性。")
    # Objective failures are valid even when other aspects still need review.
    if failures:
        result["自动_PASS_FAIL"] = "FAIL"
        result["自动_结果理由"] = "；".join(failures)
    elif answer and answer["nlg_status"] == "pass" and not reviews and intent != "unknown" and nlg:
        result["自动_PASS_FAIL"] = "PASS"
    result["自动_行为分类"] = INTENT_NAMES[intent]
    result["自动_复核状态"] = "待复核" if reviews else "已评估（建议抽检）"
    result["自动_复核原因"] = "；".join(dict.fromkeys(reviews))
    unique_issues = {(i["label"], i["reason"]) for i in issues}
    result["自动_问题标签"] = "；".join(sorted({i[0] for i in unique_issues}))
    result["自动_问题证据"] = "\n".join(label + "：" + reason for label, reason in sorted(unique_issues))
    return result


def excel_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def load_records(workbook):
    found = []
    for sheet in workbook:
        headers = [text(c.value) for c in sheet[1]]
        if "Query" not in headers:
            continue
        names = [h for h in headers if h]
        if len(names) != len(set(names)):
            raise ValueError(f"工作表{sheet.title}包含重复列名，请先修正。")
        for number in range(2, sheet.max_row + 1):
            record = {h: sheet.cell(number, i + 1).value for i, h in enumerate(headers) if h}
            if text(record.get("Query")):
                found.append((sheet.title, number, record))
    if not found:
        raise ValueError("没有找到首行包含Query的有效数据表。")
    # Group/round errors must not silently produce inaccurate history.
    seen = set()
    for sheet, number, record in found:
        group = text(record.get("会话ID"))
        round_ = text(record.get("轮次"))
        if group:
            try:
                integer = int(round_)
                if integer < 1:
                    raise ValueError()
            except ValueError:
                raise ValueError(f"{sheet}第{number}行：有会话ID时轮次须为正整数。") from None
            key = (sheet, group, integer)
            if key in seen:
                raise ValueError(f"{sheet}第{number}行：同会话轮次重复。")
            seen.add(key)
    return found


def pair_stats(records):
    groups = {}
    for sheet, number, row in records:
        group = text(row.get("配对组ID"))
        if group and not merged_rounds(row):
            groups.setdefault((sheet, group), []).append((number, row))
    stats = []
    for (sheet, group), entries in groups.items():
        if len(entries) != 2:
            stats.append({"group": group, "sheet": sheet, "note": "配对组必须恰有两条记录，本组未计算。"})
            continue
        types = {text(r.get("配对类型")) for _, r in entries}
        if len(types) != 1 or next(iter(types)) not in {"同诉求重复", "不同诉求对比"}:
            stats.append({"group": group, "sheet": sheet, "note": "请统一配对类型为同诉求重复或不同诉求对比。"})
            continue
        for field in ("music_search响应_candidates", "出歌单信息"):
            a, b = [set(song_list(r.get(field))) for _, r in entries]
            if not a or not b:
                stats.append({"group": group, "sheet": sheet, "field": field, "note": "缺少一方列表，未计算。"})
                continue
            stats.append({"group": group, "sheet": sheet, "type": next(iter(types)), "field": field,
                          "intersection": len(a & b), "jaccard": round(len(a & b) / len(a | b), 4),
                          "note": "按完整歌曲文本精确匹配；仅描述重叠，不按重叠率直接判断好坏。"})
    return stats


def report_html(items, metadata, pairs):
    counts = Counter(r["自动_PASS_FAIL"] or "待定" for _, _, _, r in items)
    total = len(items)
    known = counts["PASS"] + counts["FAIL"]
    issue_counts = Counter(label for _, _, _, r in items for label in r["自动_问题标签"].split("；") if label)
    esc = lambda v: html.escape(text(v))
    rows = []
    for sheet, number, source, result in items:
        rows.append("<tr>" + "".join("<td>" + esc(v).replace("\n", "<br>") + "</td>" for v in
                    (sheet + ":" + str(number), source["Query"], result["自动_PASS_FAIL"] or "待定",
                     result["自动_复核状态"], result["自动_问题标签"], result["自动_结果理由"],
                     result["自动_问题证据"], result["自动_复核原因"],
                     result["自动_搜索评分理由"], result["自动_播放评分理由"])) + "</tr>")
    quality_rows = []
    for field in OUTPUT_FIELDS:
        if field.endswith("分"):
            values = [r[field] for _, _, _, r in items if type(r[field]) in (int, float)]
            quality_rows.append(f"<tr><td>{esc(field)}</td><td>{len(values)}</td><td>{sum(values)/len(values):.2f}</td></tr>" if values
                                else f"<tr><td>{esc(field)}</td><td>0</td><td>未评分</td></tr>")
    rate = f"{counts['PASS']/known:.1%}" if known else "无已判定样本"
    groups = {}
    for _, _, _, r in items:
        counts_ = groups.setdefault(r["自动_行为分类"], Counter())
        counts_[r["自动_PASS_FAIL"] or "待定"] += 1
    group_rows = ''.join('<tr>' + ''.join('<td>' + esc(v) + '</td>' for v in
                         (g,c['PASS'],c['FAIL'],c['待定'])) + '</tr>' for g,c in groups.items())
    suggestions = {"调用工具错误": "优先核对搜索/播放意图与工具调度；先搜后播无需重新搜索，明确播放诉求不应再询问是否播放。",
                   "传参错误": "核对关键条件在改写中是否保留，并检查播放列表筛选依据；不要仅因空格短语而改模型。",
                   "落域错误": "核对最终标签、最终通道及混合闲聊音乐诉求的慢通道规则。"}
    action_text = ''.join('<li>' + esc(suggestions[k]) + '</li>' for k in issue_counts if k in suggestions)
    action_text += '<li>对待定、事实缺证据及ASR歧义样例补充原始NLG或带来源参考资料，再复核。</li>'
    return f"""<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>音乐Agent评估报告</title>
<style>body{{font-family:system-ui,'Microsoft YaHei';margin:32px;color:#183047}}table{{border-collapse:collapse;width:100%;margin:20px 0}}td,th{{border:1px solid #cbd5e1;padding:9px;text-align:left;vertical-align:top;overflow-wrap:anywhere}}th{{background:#edf4fa}}.summary{{padding:18px;background:#edf4fa}}pre{{white-space:pre-wrap}}button{{padding:8px}}</style>
<h1>音乐 Agent 评估报告</h1><button onclick="window.print()">打印 / 保存为PDF</button>
<p>运行ID：{esc(metadata['run_id'])}；模式：{esc(metadata['mode'])}；模型：{esc(metadata['model'])}</p>
<div class="summary">总记录 {total}；PASS {counts['PASS']}；FAIL {counts['FAIL']}；待定 {counts['待定']}。<br>
已判定样本通过率：{rate}（分母 {known}）；判定覆盖率：{known/total:.1%}。<br>
待复核记录：{sum(r['自动_复核状态']=='待复核' for _,_,_,r in items)}。待复核FAIL仍计入FAIL；请查看各条证据。</div>
<p>自动评分是辅助标注，建议抽检。PASS与链路问题独立；质量均分只统计有分数的记录，空白不是0分。ASR转写无法替代原始回复证据。具体音乐事实需参考资料核实。本次未运行真实车机，播放成功依据导出的出歌单信息。</p>
<h2>问题定位</h2><p>{esc('；'.join(f'{k}：{v}条' for k,v in issue_counts.items()) or '没有已识别标签（不代表没有问题）')}</p>
<p>标签可能同时出现，计数不能相加当作失败总数。工具可定位到字段/环节，不能证明内部代码根因。</p>
<h2>按行为分类</h2><table><tr><th>行为</th><th>PASS</th><th>FAIL</th><th>待定</th></tr>{group_rows}</table>
<h2>建议复核与改进</h2><ul>{action_text}</ul>
<h2>歌曲质量</h2><table><tr><th>维度</th><th>已评分数量</th><th>均分（0~5）</th></tr>{''.join(quality_rows)}</table>
<h2>逐条结果与改进线索</h2><table><tr><th>原表位置</th><th>Query</th><th>结果</th><th>复核状态</th><th>问题</th><th>结果理由</th><th>定位证据</th><th>复核原因</th><th>搜索质量理由</th><th>播放质量理由</th></tr>{''.join(rows)}</table>
<h2>配对重复性</h2><pre>{esc(json.dumps(pairs, ensure_ascii=False, indent=2) if pairs else '未提供配对组；未评估跨Query重复性。')}</pre>
<h2>运行记录</h2><pre>{esc(json.dumps(metadata, ensure_ascii=False, indent=2))}</pre>
</html>"""


def run(input_path, output_dir, judge=None, progress=None, cancelled=None):
    path = Path(input_path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("请使用.xlsx文件；不支持宏工作簿。")
    # Two views preserve original cells/formulas but evaluate cached exported values.
    workbook = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    records = load_records(values)
    formula_count = sum(c.data_type == "f" for s in workbook for row in s for c in row)
    if formula_count:
        raise ValueError("输入包含公式；请在Excel中将测试数据粘贴为值后重试，避免缓存数据不完整。")
    run_id = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
    metadata = {"run_id": run_id, "timestamp": datetime.now().astimezone().isoformat(),
                "mode": "模型辅助+规则检查" if judge else "仅规则检查（未完成主观评估）",
                "model": judge.config.model if judge else "未连接",
                "prompt_version": PROMPT_VERSION,
                "prompt_sha256": hashlib.sha256(SYSTEM_PROMPT.encode()).hexdigest(),
                "input_sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    evaluated = []
    for i, (sheet_name, number, row) in enumerate(records):
        if cancelled and cancelled():
            metadata["cancelled"] = True
            break
        group = text(row.get("会话ID"))
        history = sorted([r for s, _, r in records if s == sheet_name and group and
                          text(r.get("会话ID")) == group and int(text(r["轮次"])) < int(text(row["轮次"]))],
                         key=lambda r: int(text(r["轮次"])))
        result = evaluate_row(row, history, judge)
        if group and int(text(row["轮次"])) > 1:
            rounds = [int(text(r["轮次"])) for r in history]
            if rounds != list(range(1, int(text(row["轮次"])))):
                result["自动_复核状态"] = "待复核"
                result["自动_复核原因"] += "；同会话缺少前轮记录。"
                if result["自动_PASS_FAIL"] == "PASS":
                    result["自动_PASS_FAIL"] = ""
        result["自动_运行ID"] = run_id
        result["自动_评分规则版本"] = PROMPT_VERSION
        evaluated.append((sheet_name, number, row, result))
        if progress:
            progress(i + 1, len(records))
    if not evaluated:
        raise ValueError("未处理任何记录，没有生成报告。")
    metadata["processed"] = len(evaluated)
    metadata["input_records"] = len(records)
    for sheet_name in {s for s, _, _, _ in evaluated}:
        sheet = workbook[sheet_name]
        existing = {text(c.value): c.column for c in sheet[1] if c.value is not None}
        new_cols = {}
        for field in OUTPUT_FIELDS:
            # Never overwrite prior automated runs or human labels.
            header = field if field not in existing else field + "_" + run_id
            column = sheet.max_column + 1
            cell = sheet.cell(1, column, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245B78")
            sheet.column_dimensions[cell.column_letter].width = 24
            new_cols[field] = column
        for s, number, _, result in evaluated:
            if s == sheet_name:
                for field, col in new_cols.items():
                    cell = sheet.cell(number, col, excel_safe(result[field]))
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = sheet.freeze_panes or "A2"
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    target = output / ("评估结果_" + run_id + ".xlsx")
    report = output / ("评估报告_" + run_id + ".html")
    metadata_path = output / ("运行记录_" + run_id + ".json")
    workbook.save(target)
    pairs = pair_stats([(s,n,r) for s,n,r,_ in evaluated])
    report.write_text(report_html(evaluated, metadata, pairs), encoding="utf-8")
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    workbook.close()
    values.close()
    return target, report, metadata_path


def make_template(path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "测试数据"
    sheet.append(FIELDS)
    samples = [
        {"用例ID": "示例1", "会话ID": "会话A", "轮次": 1, "维度": "场景搜歌-只搜不播",
         "Query": "婚礼想放一些甜甜的歌，你有什么推荐？", "预期行为": "只搜不播", "预期需搜索": "是"},
        {"用例ID": "示例2", "会话ID": "会话A", "轮次": 2, "维度": "先搜后播",
         "Query": "就放你刚才推荐的那些吧", "预期行为": "先搜后播", "预期需搜索": "否"},
    ]
    for sample in samples:
        sheet.append([sample.get(k, "") for k in FIELDS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245B78")
        sheet.column_dimensions[cell.column_letter].width = 25
    sheet.freeze_panes = "A2"
    notes = workbook.create_sheet("填写说明")
    for line in ["示例行没有工具结果，不是真实测试，请替换后评估。", "每轮一行，同会话ID+递增轮次；本轮未调用工具则留空。",
                 "预期行为：只搜不播、又搜又播、先搜后播、其他；预期需搜索：是/否。可选列，推荐填写。",
                 "现有样例12列可直接导入；原人工标注不会发送给评分模型。",
                 "原始NLG文本优先于asr新；asr是转写文本，无法确认的差异待复核。",
                 "评测日期是原测试日期，不是自动评分日期。参考资料可包含核实的歌曲资料、事件日期、来源和摘录；仅URL不会自动联网。",
                 "时效窗口示例：测试日前14天；用于热点评估。",
                 "配对组ID须恰好两条，配对类型：同诉求重复 / 不同诉求对比。仅计算文本Jaccard重叠率，不自动给跨query质量分。",
                 "歌曲列表推荐一行一首，格式：歌名 - 歌手；保留版本。也支持JSON数组。",
                 "出歌单信息为空、[]、未找到 mediaList均代表无成功播放列表。",
                 "工具结果超过单元格容量时请在导出阶段解决；Excel单元格最多32767字符。"]:
        notes.append([line])
    notes.column_dimensions["A"].width = 115
    workbook.save(path)
